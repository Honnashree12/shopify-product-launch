import collections
import collections.abc
collections.Mapping = collections.abc.Mapping
collections.MutableMapping = collections.abc.MutableMapping
collections.Sequence = collections.abc.Sequence
collections.MutableSequence = collections.abc.MutableSequence
collections.Iterable = collections.abc.Iterable

import os
import sys
import shutil
import tempfile
import json
import httpx
from PIL import Image
import cv2
import numpy as np
from pathlib import Path
from dotenv import load_dotenv

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
load_dotenv(dotenv_path=os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".env")))

from tools.folder_tools import (
    ingest_and_extract_folder_tool,
    validate_folder_data_tool,
    create_shopify_product_tool,
    upload_folder_media_tool,
    verify_folder_product_tool,
    generate_folder_reports_tool,
)
from shopify_client import get_product, _get_config, ShopifyAPIError


class MockToolContext:
    def __init__(self, state):
        self.state = state


def delete_shopify_product(product_id):
    base_url, store_url, headers = _get_config()
    endpoint = f"{base_url}/products/{product_id}.json"
    try:
        res = httpx.delete(endpoint, headers=headers, timeout=30)
        res.raise_for_status()
        print(f"Cleaned up test product ID {product_id} from Shopify.")
    except Exception as e:
        print(f"Failed to cleanup test product {product_id}: {e}")


def create_test_case_a(path: Path):
    """
    Valid folder with multiple documents, conversions, and conflicts.
    """
    (path / "product-info.txt").write_text("Title: STEM Space Explorer Kit\nSKU: E2E-STEM-999\nBrand: ADK Labs\nCategory: Education Kits", encoding="utf-8")
    
    # Create valid DOCX
    try:
        from docx import Document
        doc = Document()
        doc.add_paragraph("This is an educational STEM Space Explorer Kit for kids. It features multiple space experiments and electronic modules.")
        doc.save(str(path / "description.docx"))
    except Exception as e:
        print(f"Failed to create DOCX: {e}")
    
    # Create CSV (Price conflict)
    (path / "pricing.csv").write_text("Field,Value\nPrice,1299\nCompare At,1999\nCost,400\n", encoding="utf-8")
    
    # price.txt (conflicting price)
    (path / "price.txt").write_text("Selling Price: 1499", encoding="utf-8")
    
    # Create XLSX for inventory locations
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Location", "Stock"])
        ws.append(["Main Warehouse", "35"])
        ws.append(["Retail Shop", "5"])
        wb.save(str(path / "inventory.xlsx"))
    except Exception as e:
        print(f"Failed to create XLSX: {e}")
        
    (path / "tags.txt").write_text("STEM, Kids, Learning, Space", encoding="utf-8")
    
    # Create Images & Videos
    images_dir = path / "images"
    images_dir.mkdir()
    
    img_hero = Image.new('RGB', (200, 200), color='red')
    img_hero.save(str(images_dir / "hero.jpg"))
    
    img_logo = Image.new('RGB', (50, 50), color='blue')
    img_logo.save(str(images_dir / "logo.png"))
    
    img_tiff = Image.new('RGB', (150, 150), color='green')
    img_tiff.save(str(images_dir / "side.tiff"))
    
    videos_dir = path / "videos"
    videos_dir.mkdir()
    try:
        fourcc = cv2.VideoWriter_fourcc(*'XVID')
        out = cv2.VideoWriter(str(videos_dir / "demo.avi"), fourcc, 20.0, (160, 120))
        for _ in range(10):
            frame = np.zeros((120, 160, 3), dtype=np.uint8)
            out.write(frame)
        out.release()
    except Exception as e:
        print(f"Failed to create video: {e}")


def main():
    print("==================================================")
    print("SHOPIFY PRODUCT CREATION AGENT END-TO-END VERIFIER")
    print("==================================================")

    # Check Environment
    store_url = os.getenv("SHOPIFY_STORE_URL")
    access_token = os.getenv("SHOPIFY_ACCESS_TOKEN")
    if not store_url or not access_token:
        print("ERROR: Shopify credentials missing in .env")
        sys.exit(1)
        
    print(f"Shopify Connected to: {store_url}")

    results = {}
    temp_dirs = []

    try:
        # 1. CREATE TEST FOLDER
        tmpdir = tempfile.mkdtemp(prefix="shopify_e2e_")
        temp_dirs.append(tmpdir)
        path = Path(tmpdir)
        create_test_case_a(path)
        print(f"Created temporary dummy folder at: {tmpdir}")

        # 2. RUN FILE DISCOVERY & AI MAPPING (Ingest)
        print("\n--- Running Ingestion Tool ---")
        state = {
            "product_folder_path": tmpdir,
            "preview_mode": True,
            "approved_to_publish": False,
        }
        context = MockToolContext(state)
        ingest_res = ingest_and_extract_folder_tool(context)
        print(ingest_res)

        # Verification of extraction
        results["File discovery"] = "PASS" if state.get("product_name") else "FAIL"
        results["Image processing (TIFF->PNG)"] = "PASS" if any("side_converted.png" in f for f in state.get("detected_images", [])) else "FAIL"
        results["Video processing (AVI->MP4)"] = "PASS" if any("demo_converted.mp4" in f for f in state.get("detected_videos", [])) else "FAIL"
        results["Field mapping"] = "PASS" if state.get("product_name") == "STEM Space Explorer Kit" else "FAIL"
        results["Conflict detection"] = "PASS" if len(state["manifest"].get("conflicts", [])) > 0 else "FAIL"

        print("\nIngested Details:")
        print(f"- Title: {state.get('product_name')}")
        print(f"- Price: {state.get('product_price')}")
        print(f"- SKU: {state.get('sku')}")
        print(f"- Conflicts: {state['manifest'].get('conflicts') if state.get('manifest') else []}")
        print(f"- Conversions: {state.get('conversions')}")

        # Fallback values for E2E validation in case of network/Gemini failure
        if not state.get("product_name"):
            print("Applying fallback values for validation and downstream testing due to Gemini API connection error...")
            state["product_name"] = "STEM Space Explorer Kit"
            state["product_description"] = "This is an educational STEM Space Explorer Kit for kids. It features multiple space experiments and electronic modules."
            state["sku"] = "E2E-STEM-999"
            state["vendor"] = "ADK Labs"
            state["tags"] = ["STEM", "Kids", "Learning", "Space"]
            state["inventory_quantity"] = 35
            state["product_status"] = "DRAFT"
            
            state["manifest"] = {
                "title": "STEM Space Explorer Kit",
                "description": "This is an educational STEM Space Explorer Kit for kids. It features multiple space experiments and electronic modules.",
                "price": 1399.0,
                "compare_at_price": 1999.0,
                "cost_per_item": 400.0,
                "category": "Education Kits",
                "product_type": "Education Kits",
                "vendor": "ADK Labs",
                "sku": "E2E-STEM-999",
                "barcode": "BARCODE999",
                "track_inventory": True,
                "inventory_quantity": 35,
                "sell_out_of_stock": False,
                "requires_shipping": True,
                "weight": 1.5,
                "weight_unit": "kg",
                "country_of_origin": "IN",
                "hs_code": "9503.00.00",
                "status": "DRAFT",
                "tags": ["STEM", "Kids", "Learning", "Space"],
                "seo_title": "STEM Space Explorer Kit",
                "seo_description": "STEM Space Explorer Kit for kids.",
                "handle": "stem-space-explorer-kit",
                "variants": [],
                "confidence": {"title": "HIGH", "price": "LOW"},
                "conflicts": [{"field": "price", "message": "pricing.csv has 1299, but price.txt has 1499"}],
                "missing_fields": []
            }

        # 3. RUN VALIDATION (Should pause for preview review)
        print("\n--- Running Validation Tool ---")
        val_res = validate_folder_data_tool(context)
        print(f"Validation Result: {val_res}")
        results["Validation pause"] = "PASS" if state.get("status") == "awaiting_approval" else "FAIL"

        # 4. RESOLVE CONFLICTS & APPROVE (Simulate Frontend Approval Submit)
        print("\n--- Resolving fields & Approving publish ---")
        state["product_price"] = 1399.0  # Resolved price (resolved conflict)
        state["approved_to_publish"] = True
        state["duplicate_action"] = "UPDATE"
        state["status"] = "running"

        # Validate again (should pass)
        val_res2 = validate_folder_data_tool(context)
        print(f"Validation Result (Post-Approval): {val_res2}")
        results["Validation resolve"] = "PASS" if "Ready for creation" in val_res2 else "FAIL"

        # 5. CREATE SHOPIFY PRODUCT
        print("\n--- Running Shopify Product Creation Tool ---")
        create_res = create_shopify_product_tool(context)
        print(create_res)
        
        prod_id = state.get("shopify_product_id")
        results["Product creation (Shopify API)"] = "PASS" if prod_id else "FAIL"
        results["Draft safety protection"] = "PASS" if state.get("product_status") == "DRAFT" else "FAIL"

        if prod_id:
            # 6. UPLOAD MEDIA
            print("\n--- Running Media Upload Tool ---")
            media_res = upload_folder_media_tool(context)
            print(media_res)
            results["Media upload (Staged uploads)"] = "PASS" if len(state.get("shopify_media", [])) > 0 else "FAIL"

            # 7. LIVE VERIFICATION
            print("\n--- Running Live Verification Tool ---")
            verify_res = verify_folder_product_tool(context)
            print(verify_res)
            results["Live verification check"] = "PASS" if "PASSED" in verify_res else "FAIL"

            # 8. GENERATE REPORT
            print("\n--- Running Report Generation Tool ---")
            report_res = generate_folder_reports_tool(context)
            print(report_res)
            results["Report generation"] = "PASS" if state.get("markdown_report_path") else "FAIL"

            # 9. FETCH FROM SHOPIFY & COMPARISON MATRIX
            print("\n--- Comparing Shopify Values ---")
            shopify_prod = get_product(prod_id)
            
            # Retrieve inventory details
            variant = shopify_prod["variants"][0] if shopify_prod.get("variants") else {}
            actual_price = float(shopify_prod.get("price", 0))
            expected_price = float(state.get("product_price", 0))
            
            # Render Comparison Table
            print("\nFIELD | EXPECTED | SHOPIFY | RESULT")
            print("-----------------------------------------")
            
            def check_field(field, expected, actual):
                match = str(expected).strip().lower() == str(actual).strip().lower()
                res = "PASS" if match else "FAIL"
                print(f"{field} | {expected} | {actual} | {res}")
                return res

            comp_title = check_field("Title", state.get("product_name"), shopify_prod.get("title"))
            comp_price = check_field("Price", expected_price, actual_price)
            comp_sku = check_field("SKU", state.get("sku"), variant.get("sku"))
            comp_vendor = check_field("Vendor", state.get("vendor"), shopify_prod.get("vendor"))
            comp_status = check_field("Status", "draft", shopify_prod.get("status"))

            results["Shopify integration match"] = "PASS" if all(x == "PASS" for x in [comp_title, comp_price, comp_sku, comp_vendor, comp_status]) else "FAIL"

            # Clean up
            delete_shopify_product(prod_id)

    except Exception as e:
        print(f"\nExecution failed: {e}")
        import traceback
        traceback.print_exc()

    finally:
        # Cleanup local directories
        for d in temp_dirs:
            try:
                shutil.rmtree(d)
            except Exception:
                pass

    # Print Final Verification Matrix
    print("\n" + "="*50)
    print("FINAL VERIFICATION MATRIX")
    print("="*50)
    for component, status in results.items():
        print(f"{component:<35} : {status}")
    print("="*50)


if __name__ == "__main__":
    main()
